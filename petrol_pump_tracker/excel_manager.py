"""
Manage Excel files for petrol pump billing - one sheet per party with running balance.
"""

import os
from datetime import datetime
from pathlib import Path

import openpyxl
from openpyxl.styles import Font, Alignment, Border, Side, PatternFill, numbers
from openpyxl.utils import get_column_letter
import pandas as pd


DATA_DIR = Path(__file__).parent / "data"
EXCEL_FILE = DATA_DIR / "petrol_pump_ledger.xlsx"

COLUMNS = [
    "S.No",
    "Date",
    "Vehicle No",
    "Fuel Type",
    "Litres",
    "Rate/Ltr",
    "Fuel Amount",
    "Cash to Driver",
    "Total Bill",
    "Payment Received",
    "Balance",
    "Notes",
]

# Column widths
COL_WIDTHS = {
    "A": 6,   # S.No
    "B": 14,  # Date
    "C": 16,  # Vehicle No
    "D": 12,  # Fuel Type
    "E": 12,  # Litres
    "F": 12,  # Rate/Ltr
    "G": 16,  # Fuel Amount
    "H": 16,  # Cash to Driver
    "I": 14,  # Total Bill
    "J": 18,  # Payment Received
    "K": 14,  # Balance
    "L": 25,  # Notes
}

HEADER_FILL = PatternFill(start_color="1F4E79", end_color="1F4E79", fill_type="solid")
HEADER_FONT = Font(name="Calibri", bold=True, color="FFFFFF", size=11)
DATA_FONT = Font(name="Calibri", size=11)
BORDER = Border(
    left=Side(style="thin"),
    right=Side(style="thin"),
    top=Side(style="thin"),
    bottom=Side(style="thin"),
)


def _ensure_workbook() -> openpyxl.Workbook:
    """Load existing workbook or create a new one."""
    DATA_DIR.mkdir(parents=True, exist_ok=True)
    if EXCEL_FILE.exists():
        return openpyxl.load_workbook(EXCEL_FILE)
    wb = openpyxl.Workbook()
    # Remove default sheet
    wb.remove(wb.active)
    return wb


def _ensure_party_sheet(wb: openpyxl.Workbook, party_name: str) -> openpyxl.worksheet.worksheet.Worksheet:
    """Get or create a sheet for a specific party with headers."""
    safe_name = party_name.strip().upper()[:31]  # Excel sheet name limit

    if safe_name in wb.sheetnames:
        return wb[safe_name]

    ws = wb.create_sheet(title=safe_name)

    # Write headers
    for col_idx, header in enumerate(COLUMNS, 1):
        cell = ws.cell(row=1, column=col_idx, value=header)
        cell.font = HEADER_FONT
        cell.fill = HEADER_FILL
        cell.alignment = Alignment(horizontal="center", vertical="center")
        cell.border = BORDER

    # Set column widths
    for col_letter, width in COL_WIDTHS.items():
        ws.column_dimensions[col_letter].width = width

    # Freeze header row
    ws.freeze_panes = "A2"

    return ws


def get_previous_balance(party_name: str) -> float:
    """Get the last balance for a party."""
    if not EXCEL_FILE.exists():
        return 0.0

    wb = openpyxl.load_workbook(EXCEL_FILE)
    safe_name = party_name.strip().upper()[:31]

    if safe_name not in wb.sheetnames:
        wb.close()
        return 0.0

    ws = wb[safe_name]
    last_row = ws.max_row

    if last_row <= 1:  # Only header
        wb.close()
        return 0.0

    balance = ws.cell(row=last_row, column=11).value  # Column K = Balance
    wb.close()
    return float(balance) if balance is not None else 0.0


def add_bill_entry(
    party_name: str,
    date: str,
    vehicle_no: str,
    fuel_type: str,
    litres: float,
    rate_per_litre: float | None,
    fuel_amount: float,
    cash_to_driver: float,
    payment_received: float = 0.0,
    notes: str = "",
) -> dict:
    """
    Add a new bill entry to the party's Excel sheet.

    Total Bill = Fuel Amount + Cash to Driver
    Balance = Previous Balance + Total Bill - Payment Received
    (Positive balance = party owes money)

    Example pink slip: Fuel 15000 + Cash to driver 2000 = Total Bill 17000

    Returns dict with entry details including new balance.
    """
    wb = _ensure_workbook()
    ws = _ensure_party_sheet(wb, party_name)

    # Find next row
    next_row = ws.max_row + 1
    serial_no = next_row - 1  # Row 1 is header

    # Calculate total bill and balance
    total_bill = fuel_amount + cash_to_driver

    prev_balance = 0.0
    if next_row > 2:  # There are previous entries
        prev_bal_cell = ws.cell(row=next_row - 1, column=11).value  # Column K = Balance
        prev_balance = float(prev_bal_cell) if prev_bal_cell is not None else 0.0

    new_balance = prev_balance + total_bill - payment_received

    # Write data
    row_data = [
        serial_no,
        date,
        vehicle_no.upper() if vehicle_no else "",
        fuel_type.upper() if fuel_type else "",
        litres,
        rate_per_litre,
        fuel_amount,
        cash_to_driver,
        total_bill,
        payment_received,
        new_balance,
        notes or "",
    ]

    for col_idx, value in enumerate(row_data, 1):
        cell = ws.cell(row=next_row, column=col_idx, value=value)
        cell.font = DATA_FONT
        cell.border = BORDER

        # Number formatting
        if col_idx in (5, 6):  # Litres, Rate
            cell.number_format = "#,##0.00"
            cell.alignment = Alignment(horizontal="right")
        elif col_idx in (7, 8, 9, 10, 11):  # Fuel Amt, Cash to Driver, Total Bill, Payment, Balance
            cell.number_format = "₹#,##0.00"
            cell.alignment = Alignment(horizontal="right")
            # Color balance: red if positive (owes), green if zero/negative
            if col_idx == 11:
                if new_balance > 0:
                    cell.font = Font(name="Calibri", size=11, color="CC0000", bold=True)
                else:
                    cell.font = Font(name="Calibri", size=11, color="006600", bold=True)
        elif col_idx == 1:  # S.No
            cell.alignment = Alignment(horizontal="center")
        elif col_idx == 2:  # Date
            cell.alignment = Alignment(horizontal="center")

    wb.save(EXCEL_FILE)
    wb.close()

    return {
        "serial_no": serial_no,
        "party_name": party_name,
        "date": date,
        "vehicle_no": vehicle_no,
        "fuel_type": fuel_type,
        "litres": litres,
        "rate_per_litre": rate_per_litre,
        "fuel_amount": fuel_amount,
        "cash_to_driver": cash_to_driver,
        "total_bill": total_bill,
        "payment_received": payment_received,
        "previous_balance": prev_balance,
        "new_balance": new_balance,
    }


def get_party_ledger(party_name: str) -> pd.DataFrame:
    """Get all entries for a party as a DataFrame."""
    if not EXCEL_FILE.exists():
        return pd.DataFrame(columns=COLUMNS)

    wb = openpyxl.load_workbook(EXCEL_FILE)
    safe_name = party_name.strip().upper()[:31]

    if safe_name not in wb.sheetnames:
        wb.close()
        return pd.DataFrame(columns=COLUMNS)

    ws = wb[safe_name]
    data = []
    for row in ws.iter_rows(min_row=2, values_only=True):
        if row[0] is not None:  # Skip empty rows
            data.append(row)

    wb.close()
    df = pd.DataFrame(data, columns=COLUMNS)
    return df


def get_all_parties() -> list[str]:
    """Get list of all party names (sheet names)."""
    if not EXCEL_FILE.exists():
        return []

    wb = openpyxl.load_workbook(EXCEL_FILE)
    parties = wb.sheetnames
    wb.close()
    return parties


def get_all_balances() -> dict[str, float]:
    """Get current balance for all parties."""
    balances = {}
    for party in get_all_parties():
        balances[party] = get_previous_balance(party)
    return balances


def delete_last_entry(party_name: str) -> bool:
    """Delete the last entry for a party (undo). Returns True if successful."""
    if not EXCEL_FILE.exists():
        return False

    wb = openpyxl.load_workbook(EXCEL_FILE)
    safe_name = party_name.strip().upper()[:31]

    if safe_name not in wb.sheetnames:
        wb.close()
        return False

    ws = wb[safe_name]
    last_row = ws.max_row

    if last_row <= 1:
        wb.close()
        return False

    ws.delete_rows(last_row)
    wb.save(EXCEL_FILE)
    wb.close()
    return True
