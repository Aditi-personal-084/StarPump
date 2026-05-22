"""Generate Excel download from database data."""

import io
from datetime import date
from typing import Optional

import openpyxl
from openpyxl.styles import Font, Alignment, Border, Side, PatternFill
import pandas as pd

from lib import models

COLUMNS = [
    "S.No", "Date", "Vehicle No", "Fuel Type", "Litres", "Rate/Ltr",
    "Fuel Amount", "Cash to Driver", "Total Bill", "Payment Received",
    "Balance", "Payment Mode", "Notes",
]

COL_WIDTHS = [6, 14, 16, 12, 12, 12, 16, 16, 14, 18, 14, 20, 25]

HEADER_FILL = PatternFill(start_color="1F4E79", end_color="1F4E79", fill_type="solid")
HEADER_FONT = Font(name="Calibri", bold=True, color="FFFFFF", size=11)
DATA_FONT = Font(name="Calibri", size=11)
BORDER = Border(
    left=Side(style="thin"), right=Side(style="thin"),
    top=Side(style="thin"), bottom=Side(style="thin"),
)


def _write_sheet(wb: openpyxl.Workbook, title: str, df: pd.DataFrame):
    """Write a single party sheet into the workbook."""
    ws = wb.create_sheet(title=title[:31])

    for col_idx, header in enumerate(COLUMNS, 1):
        cell = ws.cell(row=1, column=col_idx, value=header)
        cell.font = HEADER_FONT
        cell.fill = HEADER_FILL
        cell.alignment = Alignment(horizontal="center", vertical="center")
        cell.border = BORDER

    for col_idx, width in enumerate(COL_WIDTHS, 1):
        ws.column_dimensions[openpyxl.utils.get_column_letter(col_idx)].width = width

    ws.freeze_panes = "A2"

    if not df.empty:
        for row_idx, row in enumerate(df.itertuples(index=False), 2):
            for col_idx, value in enumerate(row, 1):
                cell = ws.cell(row=row_idx, column=col_idx, value=value)
                cell.font = DATA_FONT
                cell.border = BORDER

                if col_idx in (5, 6):
                    cell.number_format = "#,##0.00"
                    cell.alignment = Alignment(horizontal="right")
                elif col_idx in (7, 8, 9, 10, 11):
                    cell.number_format = "#,##0.00"
                    cell.alignment = Alignment(horizontal="right")
                elif col_idx in (1, 2):
                    cell.alignment = Alignment(horizontal="center")


def generate_excel() -> bytes:
    """Build an in-memory Excel workbook with one sheet per party."""
    wb = openpyxl.Workbook()
    wb.remove(wb.active)

    parties = models.get_all_parties()
    if not parties:
        ws = wb.create_sheet("No Data")
        ws.cell(row=1, column=1, value="No parties found.")
        buf = io.BytesIO()
        wb.save(buf)
        wb.close()
        return buf.getvalue()

    for party in parties:
        df = models.get_party_ledger(party)
        _write_sheet(wb, party, df)

    buf = io.BytesIO()
    wb.save(buf)
    wb.close()
    return buf.getvalue()


def generate_party_excel(
    party_name: str,
    date_from: Optional[date] = None,
    date_to: Optional[date] = None,
) -> bytes:
    """Build an Excel workbook for a single party, optionally filtered by date."""
    wb = openpyxl.Workbook()
    wb.remove(wb.active)

    df = models.get_party_ledger(party_name, date_from=date_from, date_to=date_to)
    _write_sheet(wb, party_name, df)

    buf = io.BytesIO()
    wb.save(buf)
    wb.close()
    return buf.getvalue()
